"""순위 조회 결과를 Excel 파일에 누적 저장."""
from datetime import date
from pathlib import Path
from typing import List, Optional, Tuple

from openpyxl import Workbook, load_workbook

HEADER = ["날짜", "그룹", "지점", "검색엔진", "검색어", "순위"]


def append_results(
    excel_path: Path,
    result_date: date,
    group: str,
    engine_label: str,
    results: List[Tuple[str, str, Optional[int]]],
) -> None:
    """(지점 라벨, 검색어, 순위) 목록을 엑셀 파일 맨 아래에 한 행씩 추가한다.

    파일이 없으면 헤더를 포함해 새로 생성한다. 순위가 None이면 "미노출"로 기록한다.
    """
    excel_path.parent.mkdir(parents=True, exist_ok=True)

    if excel_path.exists():
        workbook = load_workbook(excel_path)
        sheet = workbook.active
    else:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "순위기록"
        sheet.append(HEADER)

    date_str = result_date.strftime("%Y-%m-%d")
    for label, search_term, rank in results:
        sheet.append(
            [date_str, group, label, engine_label, search_term, rank if rank is not None else "미노출"]
        )

    workbook.save(excel_path)
